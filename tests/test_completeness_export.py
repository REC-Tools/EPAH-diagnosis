import csv
import importlib.util
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("completeness_export", ROOT / "scripts/export_completeness.py")
module = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
SPEC.loader.exec_module(module)


def test_excel_and_csv_report(tmp_path):
    generate = module.load_generate()
    report = generate(ROOT / "data/verification_matrix.csv",
                      ROOT / "data/municipalities.csv", ROOT / "config/indicators.json")
    contents = module.tables(report)
    assert len(contents["Municipalities"]) == 28
    assert len(contents["Indicators"]) == 10
    for name, rows in contents.items():
        module.write_csv(tmp_path / f"{name}.csv", rows)
        with (tmp_path / f"{name}.csv").open(encoding="utf-8-sig", newline="") as f:
            assert len(list(csv.reader(f))) == len(rows)
    workbook = tmp_path / "completeness.xlsx"
    module.write_excel(workbook, contents)
    wb = load_workbook(workbook, read_only=True, data_only=True)
    assert wb.sheetnames == ["Summary", "Municipalities", "Indicators"]
    assert wb["Municipalities"].max_row == 28
    assert wb["Indicators"].max_row == 10
    assert wb["Summary"]["B4"].value == 243
