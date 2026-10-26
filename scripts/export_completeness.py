"""Export the offline completeness inventory to CSV and formatted XLSX."""
from __future__ import annotations

import argparse
import csv
import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def load_generate():
    source = Path(__file__).with_name("report_completeness.py")
    spec = importlib.util.spec_from_file_location("report_completeness", source)
    if not spec or not spec.loader:
        raise RuntimeError("Unable to load completeness reporter")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.generate


def tables(report: dict) -> dict[str, list[list[object]]]:
    indicators = [["indicator_id", "role", "validated", "acquired", "to_verify",
                   "other_status", "nonempty_values"]]
    for row in report["by_indicator"]:
        s = row["status_counts"]
        indicators.append([row["indicator_id"], row["role"], s.get("Dato validato", 0),
                           s.get("Dato acquisito", 0), s.get("Da verificare", 0),
                           sum(v for k, v in s.items() if k not in
                               ("Dato validato", "Dato acquisito", "Da verificare")),
                           row["nonempty_values"]])
    municipalities = [["municipality", "validated", "acquired", "missing_or_unvalidated",
                       "nonempty_values"]]
    municipalities += [[r["municipality"], r["validated"], r["acquired"],
                        r["missing_or_unvalidated"], r["nonempty_values"]]
                       for r in report["by_municipality"]]
    summary = [["metric", "value"], ["municipalities", report["municipalities"]],
               ["indicators", report["indicators"]], ["checks", report["checks"]]]
    summary += [[f"status: {k}", v] for k, v in report["status_counts"].items()]
    return {"Summary": summary, "Municipalities": municipalities, "Indicators": indicators}


def write_csv(path: Path, rows: list[list[object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerows(rows)


def write_excel(path: Path, contents: dict[str, list[list[object]]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in contents.items():
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[letter].width = min(42, max(14, width + 3))
    wb.save(path)


def main() -> None:
    p = argparse.ArgumentParser(description="Export completeness inventory to Excel and CSV")
    p.add_argument("--matrix", type=Path, default=Path("data/verification_matrix.csv"))
    p.add_argument("--municipalities", type=Path, default=Path("data/municipalities.csv"))
    p.add_argument("--catalog", type=Path, default=Path("config/indicators.json"))
    p.add_argument("--output-dir", type=Path, default=Path("outputs/quality"))
    args = p.parse_args()
    report = load_generate()(args.matrix, args.municipalities, args.catalog)
    contents = tables(report)
    targets = [args.output_dir / f"completeness_{name.lower()}.csv" for name in contents]
    targets.append(args.output_dir / "completeness.xlsx")
    sources = {args.matrix.resolve(), args.municipalities.resolve(), args.catalog.resolve()}
    if any(target.resolve() in sources for target in targets):
        p.error("Export path would overwrite an input")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name, rows in contents.items():
        write_csv(args.output_dir / f"completeness_{name.lower()}.csv", rows)
    write_excel(args.output_dir / "completeness.xlsx", contents)
    print(f"Exported {report['checks']} checks to {args.output_dir}")


if __name__ == "__main__":
    main()
