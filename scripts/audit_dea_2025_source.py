"""Read the Ministry of Health 2025 emergency network XLSX without promoting data.

Creates a candidate CSV and provenance manifest in a new, ignored output folder.
It does not validate geographic access or certify independent scientific review.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_URL = "https://www.salute.gov.it/new/it/banche-dati/elenco-strutture-della-rete-dellemergenza-ospedaliera/"
FIELDS = ("hospital_name", "municipality", "dea_level", "reference_year", "source")
REQUIRED = ("Anno", "Codice Regione", "Codice struttura", "Subcodice", "Denominazione struttura", "Comune", "Livello Emergenza")


def extract(source: Path) -> tuple[list[dict[str, str]], dict]:
    wb = load_workbook(source, read_only=True, data_only=True)
    if "Strutture_Emergenza_2025" not in wb.sheetnames:
        raise ValueError("Expected Ministry worksheet Strutture_Emergenza_2025")
    sheet = wb["Strutture_Emergenza_2025"]
    sheet.reset_dimensions()
    iterator = (row for row in sheet.iter_rows(values_only=True) if row and row[0] is not None)
    headers = [str(value).strip() for value in next(iterator)]
    if len(set(headers)) != len(headers) or not set(REQUIRED) <= set(headers):
        raise ValueError("Missing or duplicated Ministry header columns")
    records = [dict(zip(headers, row)) for row in iterator]
    if not records or any(str(row["Anno"]).strip() != "2025" for row in records):
        raise ValueError("Expected exclusively 2025 records")
    piem = [row for row in records if str(row["Codice Regione"]).strip() == "010"]
    types = Counter(str(row["Livello Emergenza"] or "").strip() for row in piem)
    eligible = [row for row in piem if str(row["Livello Emergenza"]).strip() in {"DEA1", "DEA2"}]
    candidates = []
    facility_keys: set[tuple[str, str]] = set()
    for row in eligible:
        name = str(row["Denominazione struttura"] or "").strip()
        municipality = str(row["Comune"] or "").strip().strip("'").title()
        facility = str(row["Codice struttura"] or "").strip()
        subcode = str(row["Subcodice"] or "").strip()
        key = (facility, subcode)
        if not name or not municipality or len(facility) != 6 or not facility.isdigit() or key in facility_keys:
            raise ValueError(f"Empty, invalid or duplicated Ministry facility key: {key}")
        facility_keys.add(key)
        candidates.append({
            "hospital_name": name,
            "municipality": municipality,
            "dea_level": "I" if row["Livello Emergenza"] == "DEA1" else "II",
            "reference_year": "2025",
            "source": "Ministero della Salute - Strutture Emergenza 2025; codice struttura "
                      + facility + "; subcodice " + (subcode or "-"),
        })
    report = {
        "source_url": SOURCE_URL,
        "source_filename": source.name,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "reference_year": "2025",
        "national_rows": len(records),
        "piemonte_rows": len(piem),
        "piemonte_levels": dict(sorted(types.items())),
        "dea_facilities": len(candidates),
        "dea_municipalities": len({item["municipality"] for item in candidates}),
        "reviewed_by_human": False,
        "validated_for_release": False,
        "route_times_recomputed": False,
    }
    wb.close()
    return candidates, report


def candidate_csv(rows: list[dict[str, str]]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=FIELDS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8")


def compare_candidate_reference(candidate: bytes, reference: bytes) -> dict:
    """Accept Git worktree CRLF conversion only, never changed CSV content.

    The generated candidate always uses LF. The committed reference may have
    CRLF endings after Git for Windows checks out a text CSV.
    """
    exact = reference == candidate
    normalized = reference.replace(b"\r\n", b"\n")
    match = normalized == candidate
    return {
        "matches_reference": match,
        "matches_reference_bytes": exact,
        "reference_line_endings_only_difference": match and not exact,
        "reference_sha256": hashlib.sha256(reference).hexdigest(),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare official DEA 2025 XLSX to curated Piemonte reference (read-only)")
    parser.add_argument("--source-xlsx", type=Path, required=True)
    parser.add_argument("--reference", type=Path, default=ROOT / "data/reference/piemonte_dea_hospitals.csv")
    parser.add_argument("--out-dir", type=Path, default=ROOT / "outputs/quality/dea_2025_ministry_audit")
    parser.add_argument("--strict", action="store_true", help="Fail if candidate and reference differ")
    args = parser.parse_args()
    if args.out_dir.exists():
        raise FileExistsError("Use a fresh --out-dir; existing evidence must not be overwritten")
    rows, report = extract(args.source_xlsx)
    payload = candidate_csv(rows)
    report["candidate_sha256"] = hashlib.sha256(payload).hexdigest()
    report.update(compare_candidate_reference(payload, args.reference.read_bytes()))
    report["reference_path"] = str(args.reference)
    report["review_needed"] = True
    args.out_dir.mkdir(parents=True)
    (args.out_dir / "dea_hospitals_2025_candidate.csv").write_bytes(payload)
    (args.out_dir / "audit.json").write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if args.strict and not report["matches_reference"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
