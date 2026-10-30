"""Read-only 2025 national DEA alternatives and original ISTAT route samples."""
import argparse
import csv
import hashlib
import io
import json
import math
import zipfile
from pathlib import Path

import shapefile
from openpyxl import load_workbook
from epah_diagnosis.service_access import _as_float, _detect_delimiter, _matrix_columns, normalize_text

SAMPLE_ORIGINS = ("001073", "001224", "001066")
# Verified against ISTAT 2026 codes, but 2021 continuity requires human review.
ALIASES = {
    (4, "bolzano"): ("021008", "bolzanobozen"),
    (5, "negrar"): ("023052", "negrardivalpolicella"),
    (13, "popoli"): ("068033", "popoliterme"),
}


def digest(path):
    hasher = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1048576), b""):
            hasher.update(block)
    return hasher.hexdigest()


def code(value):
    txt = str(value or "").strip()
    return txt.zfill(6) if txt.isdigit() and len(txt) <= 6 else ""


def boundaries(archive_path):
    with zipfile.ZipFile(archive_path) as archive:
        names = [x for x in archive.namelist() if x.lower().endswith(".dbf")
                 and Path(x).name.lower().startswith("com0101")]
        if len(names) != 1:
            raise ValueError("Expected exactly one official ISTAT municipality DBF")
        reader = shapefile.Reader(dbf=io.BytesIO(archive.read(names[0])))
        result = {}
        for entry in reader.iterRecords():
            item = entry.as_dict()
            key = int(item["COD_REG"]), normalize_text(item["COMUNE"])
            value = code(item["PRO_COM_T"])
            if not key[1] or not value:
                raise ValueError("Invalid official municipality code")
            result.setdefault(key, set()).add(value)
    return result


def national_dea(xlsx, official):
    book = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        sheet = book["Strutture_Emergenza_2025"]
        sheet.reset_dimensions()
        iterator = (r for r in sheet.iter_rows(values_only=True) if r and r[0] is not None)
        header = [str(x).strip() for x in next(iterator)]
        needed = {"Anno", "Codice Regione", "Comune", "Livello Emergenza", "Denominazione struttura"}
        if not needed <= set(header) or len(header) != len(set(header)):
            raise ValueError("Invalid Ministry columns")
        data = [dict(zip(header, row)) for row in iterator]
    finally:
        book.close()
    if not data or any(str(r["Anno"]).strip() != "2025" for r in data):
        raise ValueError("Expected official Ministry year 2025")
    eligible = [r for r in data if str(r["Livello Emergenza"]).strip() in {"DEA1", "DEA2"}]
    locations, unresolved, aliases = {}, [], []
    for row in eligible:
        src_region = str(row["Codice Regione"]).strip()
        if not src_region.isdigit():
            raise ValueError("Invalid ministry region")
        num = int(src_region)
        region = 4 if num in (41, 42) else num // 10 if num % 10 == 0 else None
        if region is None:
            raise ValueError("Unsupported Ministry region code")
        name = str(row["Comune"] or "").strip().strip("'")
        key = region, normalize_text(name)
        found = official.get(key, set())
        if not found and key in ALIASES:
            istat, target_name = ALIASES[key]
            if official.get((region, target_name), set()) != {istat}:
                raise ValueError("Undocumented hospital municipality alias: " + name)
            found = {istat}
            aliases.append({"ministry_name": name, "istat_2026_code": istat,
                            "historical_2021_crosswalk_review_required": True})
        if len(found) != 1:
            unresolved.append({"region": src_region, "name": name, "matched_codes": sorted(found)})
            continue
        istat = next(iter(found))
        locations.setdefault(istat, {"municipality": name, "region": src_region})
    stats = {"national_rows": len(data), "national_dea_facilities": len(eligible),
             "distinct_dea_municipalities": len(locations),
             "piemonte_dea_municipalities": sum(v["region"] == "010" for v in locations.values()),
             "outside_dea_municipalities": sum(v["region"] != "010" for v in locations.values()),
             "historical_name_aliases_to_review": aliases}
    return locations, unresolved, stats


def read_road(archive_path, origins, target_codes):
    """One original-CSV pass; retain source ZIP member and row for human audit."""
    found, valid_schema = {}, False
    with zipfile.ZipFile(archive_path) as archive:
        names = sorted(n for n in archive.namelist() if n.lower().endswith(".csv"))
        if not names:
            raise ValueError("Missing ISTAT road CSV members")
        for name in names:
            with archive.open(name) as raw:
                stream = io.TextIOWrapper(raw, encoding="utf-8-sig", errors="replace", newline="")
                delim = _detect_delimiter(stream.read(8192))
                stream.seek(0)
                reader = csv.DictReader(stream, delimiter=delim)
                if not reader.fieldnames:
                    continue
                try:
                    fields = _matrix_columns(reader.fieldnames)
                except ValueError:
                    continue
                if "distance" not in fields:
                    raise ValueError("Missing required original KM_TOT column")
                valid_schema = True
                for line, row in enumerate(reader, 2):
                    origin = code(row.get(fields["origin_code"]))
                    if origin not in origins:
                        continue
                    dest = code(row.get(fields["destination_code"]))
                    if dest not in target_codes:
                        continue
                    key = origin, dest
                    if key in found:
                        raise ValueError("Duplicated official road pair: " + str(key))
                    minutes, km = _as_float(row[fields["time"]]), _as_float(row[fields["distance"]])
                    if not all(map(math.isfinite, (minutes, km))) or not (0 <= minutes <= 1440 and 0 <= km <= 3000):
                        raise ValueError("Invalid official road time/distance: " + str(key))
                    found[key] = (minutes, km, name, line)
    if not valid_schema:
        raise ValueError("No usable original ISTAT CSV schema")
    return found


def audit(road, bounds, ministry, draft):
    qa = json.loads((draft / "dea_accessibility_qa.json").read_text(encoding="utf-8"))
    data_file = draft / "dea_accessibility_minutes.csv"
    for field, source in (("road_zip_sha256", road), ("boundaries_zip_sha256", bounds),
                          ("data_csv_sha256", data_file)):
        if qa.get(field) != digest(source):
            raise ValueError("Original source does not match QA SHA-256: " + field)
    if qa.get("qa_status") != "technical_checks_only" or qa.get("validated_for_release") is not False:
        raise ValueError("Expected unapproved draft")
    with data_file.open(encoding="utf-8-sig", newline="") as f:
        observations = list(csv.DictReader(f))
    rows = {r["municipality_code"]: r for r in observations}
    if len(rows) != 27 or len(observations) != 27 or set(SAMPLE_ORIGINS) - set(rows):
        raise ValueError("Wrong 27-municipality coverage")
    hospitals, missing, source_stats = national_dea(ministry, boundaries(bounds))
    outside = {k for k, v in hospitals.items() if v["region"] != "010"}
    samples = {(o, rows[o]["nearest_dea_code"]) for o in SAMPLE_ORIGINS}
    routes = read_road(road, set(rows), outside | {d for _, d in samples})
    sample_checks = []
    for origin, destination in sorted(samples):
        value = routes.get((origin, destination))
        row = rows[origin]
        good = bool(value and abs(value[0] - float(row["travel_time_minutes"])) <= .011
                    and abs(value[1] - float(row["distance_km"])) <= .011)
        sample_checks.append({"origin": origin, "destination": destination,
                              "expected_minutes": row["travel_time_minutes"],
                              "expected_km": row["distance_km"],
                              "original_minutes": value[0] if value else None,
                              "original_km": value[1] if value else None,
                              "zip_member": value[2] if value else None,
                              "csv_row_number": value[3] if value else None,
                              "matches_ISTAT_original": good})
    outside_found = {(o, d): r for (o, d), r in routes.items() if d in outside}
    by_origin = {}
    for (origin, dest), val in outside_found.items():
        candidate = val[0], dest, val[1]
        if origin not in by_origin or candidate < by_origin[origin]:
            by_origin[origin] = candidate
    comparison = []
    for origin, row in sorted(rows.items()):
        nearest = by_origin.get(origin)
        current = float(row["travel_time_minutes"])
        comparison.append({"municipality_code": origin, "municipality": row["municipality"],
                           "piemonte_minutes": current,
                           "nearest_outside_code": nearest[1] if nearest else None,
                           "nearest_outside_name": hospitals[nearest[1]]["municipality"] if nearest else None,
                           "outside_minutes": nearest[0] if nearest else None,
                           "outside_faster": bool(nearest and nearest[0] < current)})
    faster = sum(x["outside_faster"] for x in comparison)
    expected = len(outside) * len(rows)
    blockers = []
    if missing:
        blockers.append("Unresolved hospital place names: " + str(len(missing)))
    if len(outside_found) != expected:
        blockers.append("Incomplete national road pairs: %d/%d" % (len(outside_found), expected))
    if not all(x["matches_ISTAT_original"] for x in sample_checks):
        blockers.append("Three original ISTAT sample routes not all matched")
    if faster:
        blockers.append("National DEA closer than Piemonte-only result for %d origins" % faster)
    return {"indicator_id": "dea_accessibility", "source_road_sha256": digest(road),
            "source_boundaries_sha256": digest(bounds), "source_ministry_sha256": digest(ministry),
            "draft_data_sha256": qa["data_csv_sha256"],
            "ministry_2025": source_stats, "unresolved": missing,
            "outside_pairs_observed": len(outside_found), "outside_pairs_expected": expected,
            "original_route_samples": sample_checks, "municipality_comparison": comparison,
            "origins_with_faster_outside_dea": faster,
            "technical_blockers": blockers, "technical_pass": not blockers,
            "validated_for_release": False, "human_review_required": True,
            "caveat": "2021 municipality geography, 2020 road graph, 2025 hospital snapshot; 2021/2026 historical code review still required. Not real ambulance access."}


def main():
    p = argparse.ArgumentParser(description="Read-only DEA 2025 national source & route sample audit")
    for name in ("road-zip", "boundaries-zip", "ministry-xlsx", "draft-dir", "out-dir"):
        p.add_argument("--" + name, required=True, type=Path)
    p.add_argument("--strict", action="store_true")
    a = p.parse_args()
    if a.out_dir.exists():
        raise FileExistsError("Use a fresh --out-dir: evidence must not be overwritten")
    result = audit(a.road_zip, a.boundaries_zip, a.ministry_xlsx, a.draft_dir)
    a.out_dir.mkdir(parents=True)
    (a.out_dir / "dea_national_coverage_audit.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("Ministero 2025:", result["ministry_2025"])
    print("Outside route pairs:", result["outside_pairs_observed"], "/", result["outside_pairs_expected"])
    print("Original ISTAT samples pass:", [x["matches_ISTAT_original"] for x in result["original_route_samples"]])
    print("Faster DEA outside Piemonte:", result["origins_with_faster_outside_dea"])
    print("Technical audit:", "PASS" if result["technical_pass"] else "BLOCKED")
    for blocker in result["technical_blockers"]:
        print("  -", blocker)
    print("No canonical data or approvals changed.")
    if a.strict and result["technical_blockers"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
