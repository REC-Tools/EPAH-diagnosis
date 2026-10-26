from __future__ import annotations

import io
import re
import unicodedata
import urllib.request
from dataclasses import dataclass

from openpyxl import load_workbook


SNAI_2020_URL = (
    "https://politichecoesione.governo.it/media/2825/"
    "20220214-mappa-ai-2020-elenco-e-classificazione-comuni.xlsx"
)

CLASS_LABELS = {
    "A": "Polo",
    "B": "Polo intercomunale",
    "C": "Cintura",
    "D": "Intermedio",
    "E": "Periferico",
    "F": "Ultraperiferico",
}


@dataclass(frozen=True)
class SnaiRecord:
    municipality: str
    snai_class: str
    snai_label: str


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", text.casefold())


def _find_columns(rows: list[tuple]) -> tuple[int, int, int]:
    for idx, row in enumerate(rows[:40]):
        labels = [normalize_text(value) for value in row]

        name_idx = next(
            (
                i for i, value in enumerate(labels)
                if value in {
                    "comune",
                    "denominazionecomune",
                    "nomecomune",
                    "denominazionedelcomune",
                    "comune30sett20",
                }
                or "comune30sett20" in value
                or (
                    "comune" in value
                    and "cod" not in value
                    and "numero" not in value
                )
            ),
            None,
        )

        class_idx = next(
            (
                i for i, value in enumerate(labels)
                if value in {
                    "classe",
                    "classeareeinterne",
                    "classificazione",
                    "classificazionesnai",
                    "classificazioneareeinterne",
                    "codai2020",
                }
                or "codai2020" in value
                or ("classe" in value and "comune" not in value)
            ),
            None,
        )

        if name_idx is not None and class_idx is not None:
            return idx, name_idx, class_idx

    preview = []
    for idx, row in enumerate(rows[:10], start=1):
        values = " | ".join(str(v) for v in row if v not in (None, ""))
        if values:
            preview.append(f"row {idx}: {values[:300]}")

    raise ValueError(
        "Unable to identify municipality and SNAI-class columns. "
        + "; ".join(preview)
    )


def _extract_class(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip().upper()

    direct = re.fullmatch(r"[A-F]", text)
    if direct:
        return text

    match = re.search(r"\b([A-F])\b", text)
    if match:
        return match.group(1)

    normalized = normalize_text(text)
    for letter, label in CLASS_LABELS.items():
        if normalize_text(label) == normalized:
            return letter

    return None


def parse_snai_workbook(payload: bytes) -> list[SnaiRecord]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    errors: list[str] = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        try:
            header_idx, name_idx, class_idx = _find_columns(rows)
        except ValueError as exc:
            errors.append(f"{sheet.title}: {exc}")
            continue

        output: list[SnaiRecord] = []
        for row in rows[header_idx + 1:]:
            if name_idx >= len(row) or class_idx >= len(row):
                continue

            municipality = row[name_idx]
            snai_class = _extract_class(row[class_idx])
            if municipality in (None, "") or snai_class is None:
                continue

            output.append(
                SnaiRecord(
                    municipality=str(municipality).strip(),
                    snai_class=snai_class,
                    snai_label=CLASS_LABELS[snai_class],
                )
            )

        if output:
            return output

    raise ValueError(
        "No usable SNAI classification table found. "
        + " | ".join(errors[:5])
    )


def download_snai(url: str = SNAI_2020_URL) -> list[SnaiRecord]:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 EPAH-diagnosis/1.0"},
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        payload = response.read()
    return parse_snai_workbook(payload)


def filter_municipalities(
    records: list[SnaiRecord],
    municipality_names: set[str],
) -> list[SnaiRecord]:
    wanted = {normalize_text(name): name for name in municipality_names}
    selected = [
        record
        for record in records
        if normalize_text(record.municipality) in wanted
    ]

    found = {normalize_text(record.municipality) for record in selected}
    missing = sorted(
        original for key, original in wanted.items() if key not in found
    )
    if missing:
        raise ValueError(
            "Municipalities not found in SNAI dataset: " + ", ".join(missing)
        )

    if len(selected) != len(wanted):
        raise ValueError(
            f"Expected {len(wanted)} municipalities, found {len(selected)} rows."
        )

    return selected
