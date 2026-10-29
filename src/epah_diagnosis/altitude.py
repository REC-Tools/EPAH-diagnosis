from __future__ import annotations

import io
import re
import unicodedata
import urllib.request
from dataclasses import dataclass

from openpyxl import load_workbook


ALTITUDE_XLSX_URL = (
    "https://www.istat.it/wp-content/uploads/2026/02/"
    "Altimetria_Comuni-al-31_12_2021.xlsx"
)


@dataclass(frozen=True)
class AltitudeRecord:
    municipality: str
    mean_altitude_m: float
    min_altitude_m: float | None = None
    max_altitude_m: float | None = None
    municipality_code: str | None = None




def parse_istat_code(value: object) -> str | None:
    """Normalize official six-digit PRO_COM/PROCOM_T codes, retaining leading zeros."""
    if isinstance(value, float):
        if not value.is_integer():
            return None
        value = int(value)
    raw = "" if value is None else str(value).strip()
    if not raw.isdigit() or not 1 <= len(raw) <= 6:
        return None
    return raw.zfill(6)

def normalize_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", text.casefold())


def _find_header_row(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    max_scan = min(30, len(rows))

    # First prefer the real ISTAT one-row schema. This avoids accidentally
    # combining a title row with the actual header and shifting column meaning.
    for idx in range(max_scan):
        labels = [normalize_text(value) for value in rows[idx]]

        if "comune" not in labels or "media" not in labels:
            continue

        name_idx = labels.index("comune")
        mean_idx = labels.index("media")
        min_idx = labels.index("altmin") if "altmin" in labels else -1
        max_idx = labels.index("altmax") if "altmax" in labels else -1

        return idx, {
            "municipality": name_idx,
            "mean": mean_idx,
            "min": min_idx,
            "max": max_idx,
            "code": next((i for i, label in enumerate(labels) if label in {"procom", "procomt", "codicecomune"}), -1),
        }

    # Fallback for alternative / multi-row workbook layouts.
    for idx in range(max_scan):
        for depth in (1, 2, 3):
            if idx + depth > max_scan:
                continue

            width = max(len(rows[row_idx]) for row_idx in range(idx, idx + depth))
            labels: list[str] = []

            for col_idx in range(width):
                parts = []
                for row_idx in range(idx, idx + depth):
                    row = rows[row_idx]
                    value = row[col_idx] if col_idx < len(row) else None
                    normalized = normalize_text(value)
                    if normalized:
                        parts.append(normalized)
                labels.append("".join(parts))

            name_idx = next(
                (
                    i for i, value in enumerate(labels)
                    if value in {
                        "comune",
                        "denominazionecomune",
                        "nomecomune",
                        "denominazionedelcomune",
                    }
                    or (
                        "comune" in value
                        and "cod" not in value
                        and "numero" not in value
                    )
                ),
                None,
            )

            mean_idx = next(
                (
                    i for i, value in enumerate(labels)
                    if value in {
                        "altitudinemedia",
                        "quotamedia",
                        "altmedia",
                        "mediaaltitudine",
                        "valoremedio",
                    }
                    or (
                        "valoremedio" in value
                        and "valoremediano" not in value
                    )
                    or (
                        "media" in value
                        and "mediana" not in value
                        and ("alt" in value or "quota" in value)
                    )
                ),
                None,
            )

            if name_idx is None or mean_idx is None:
                continue

            min_idx = next(
                (
                    i for i, value in enumerate(labels)
                    if value in {
                        "altitudineminima",
                        "quotaminima",
                        "altmin",
                        "valoreminimo",
                    }
                    or "valoreminimo" in value
                    or (
                        ("min" in value or "minima" in value)
                        and ("alt" in value or "quota" in value)
                    )
                ),
                None,
            )

            max_idx = next(
                (
                    i for i, value in enumerate(labels)
                    if value in {
                        "altitudinemassima",
                        "quotamassima",
                        "altmax",
                        "valoremassimo",
                    }
                    or "valoremassimo" in value
                    or (
                        ("max" in value or "massima" in value)
                        and ("alt" in value or "quota" in value)
                    )
                ),
                None,
            )

            return idx + depth - 1, {
                "municipality": name_idx,
                "mean": mean_idx,
                "min": min_idx if min_idx is not None else -1,
                "max": max_idx if max_idx is not None else -1,
                "code": next((i for i, label in enumerate(labels) if label in {"procom", "procomt", "codicecomune"}), -1),
            }

    preview_rows = []
    for idx, row in enumerate(rows[:8], start=1):
        preview = " | ".join(str(value) for value in row if value not in (None, ""))
        if preview:
            preview_rows.append(f"row {idx}: {preview[:300]}")

    details = "; ".join(preview_rows)
    raise ValueError(
        "Unable to identify municipality and mean-altitude columns."
        + (f" Header preview: {details}" if details else "")
    )

def _as_float(value: object) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(".", "").replace(",", ".")
    return float(text)


def parse_altitude_workbook(payload: bytes) -> list[AltitudeRecord]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)

    errors: list[str] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        try:
            header_idx, columns = _find_header_row(rows)
        except ValueError as exc:
            errors.append(f"{sheet.title}: {exc}")
            continue

        output: list[AltitudeRecord] = []
        for row in rows[header_idx + 1:]:
            if not row:
                continue

            if columns["municipality"] >= len(row) or columns["mean"] >= len(row):
                continue

            name = row[columns["municipality"]]
            mean = row[columns["mean"]]
            if name in (None, "") or mean in (None, ""):
                continue

            try:
                mean_value = _as_float(mean)
            except (TypeError, ValueError):
                # ISTAT includes descriptive / metadata rows below the
                # multi-line header. These are not municipality observations.
                continue

            if mean_value is None:
                continue

            def optional_numeric(column_index: int) -> float | None:
                if column_index < 0 or column_index >= len(row):
                    return None
                try:
                    return _as_float(row[column_index])
                except (TypeError, ValueError):
                    return None

            min_value = optional_numeric(columns["min"])
            max_value = optional_numeric(columns["max"])

            output.append(
                AltitudeRecord(
                    municipality=str(name).strip(),
                    mean_altitude_m=float(mean_value),
                    min_altitude_m=min_value,
                    max_altitude_m=max_value,
                    municipality_code=(
                        parse_istat_code(row[columns["code"]])
                        if columns["code"] >= 0 and columns["code"] < len(row)
                        else None
                    ),
                )
            )

        if output:
            return output

    raise ValueError(
        "No usable altitude table found in ISTAT workbook. "
        + " | ".join(errors[:5])
    )


def download_altitude(url: str = ALTITUDE_XLSX_URL) -> list[AltitudeRecord]:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 EPAH-diagnosis/1.0"},
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        payload = response.read()
    return parse_altitude_workbook(payload)


def filter_municipalities(
    records: list[AltitudeRecord],
    municipality_names: set[str],
) -> list[AltitudeRecord]:
    wanted = {normalize_text(name): name for name in municipality_names}
    selected = [
        record
        for record in records
        if normalize_text(record.municipality) in wanted
    ]

    found = {normalize_text(record.municipality) for record in selected}
    missing = sorted(
        original for key, original in wanted.items() if key not in found
    )
    if missing:
        raise ValueError(
            "Municipalities not found in ISTAT altitude dataset: "
            + ", ".join(missing)
        )

    if len(selected) != len(wanted):
        raise ValueError(
            f"Expected {len(wanted)} municipalities, found {len(selected)} rows."
        )
    return selected
