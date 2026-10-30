"""Strict structure checks and finished v0.1 diagnostic artifacts.

Nothing in this module changes verification status or certifies source validity.
Call only after the input release gate has accepted real reviewed observations.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import struct
from collections import Counter
from html import escape
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from epah_diagnosis.cli import read_municipality_codes
from epah_diagnosis.mapping import NUMERIC_FIELDS
from epah_diagnosis.cartography import (
    _oriented_polygon_path, _polygons, _ring_area, render_thematic_map,
)
from epah_diagnosis.municipality_profile import NUMERIC_INDICATORS, CATEGORICAL_INDICATORS

SNAI_COLORS = {
    "A": "#2b8cbe", "B": "#7bccc4", "C": "#bae4bc",
    "D": "#fdd49e", "E": "#fc8d59", "F": "#b30000",
}
MAP_SPECS = (
    ("population_65_plus_pct", "Popolazione con almeno 65 anni", "%", "map_population_65_plus.png"),
    ("income_vulnerability_pct", "Contribuenti con reddito <= 15.000 EUR", "%", "map_income_vulnerability.png"),
    ("old_building_stock_pct", "Edifici residenziali prima del 1981", "%", "map_old_building_stock.png"),
    ("heating_degree_days", "Gradi-giorno riscaldamento", "GG", "map_heating_degree_days.png"),
    ("mean_altitude_m", "Altitudine media comunale", "m s.l.m.", "map_mean_altitude.png"),
    ("dea_access_minutes", "Accesso al DEA I/II (proxy)", "minuti", "map_dea_access_minutes.png"),
    ("upper_secondary_access_minutes", "Accesso ai poli scolastici superiori", "minuti", "map_upper_secondary_access_minutes.png"),
    ("rail_access_minutes", "Accesso ai gateway ferroviari", "minuti", "map_rail_access_minutes.png"),
    ("snai_class", "Classificazione SNAI 2020", "classi A-F", "map_snai_classification.png"),
)
EXCLUDED_HASH_FILES = {"release_manifest.json", "readiness.json"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1048576), b""):
            digest.update(chunk)
    return digest.hexdigest()


def check_geojson(data: dict, registry: dict[str, str], *, enriched=False) -> None:
    if data.get("type") != "FeatureCollection":
        raise ValueError("Official geometry must be a GeoJSON FeatureCollection")
    features = data.get("features")
    if not isinstance(features, list) or len(features) != 27:
        raise ValueError("Exactly 27 official municipality features required")
    names = set()
    codes = set()
    for feature in features:
        if feature.get("type") != "Feature":
            raise ValueError("Invalid feature type")
        props = feature.get("properties") or {}
        name, code = props.get("municipality"), props.get("municipality_code")
        if name not in registry or code != registry[name]:
            raise ValueError(f"Unmatched municipality name/code in official geography: {name!r}/{code!r}")
        if name in names or code in codes:
            raise ValueError("Duplicate municipality feature or ISTAT code")
        names.add(name)
        codes.add(code)
        for _ in _polygons(feature.get("geometry") or {}):
            pass
        if enriched:
            for field in NUMERIC_FIELDS:
                val = props.get(field)
                if (not isinstance(val, (int, float)) or
                    not math.isfinite(float(val))):
                    raise ValueError(f"Nonfinite/missing numeric geo property: {name}/{field}")
            if props.get("snai_class") not in SNAI_COLORS:
                raise ValueError(f"Invalid categorical SNAI property for {name}")
    if names != set(registry):
        raise ValueError("27 geographic features do not match canonical registry")


# These are vulnerability/pressure directions, not causal energy-poverty
# judgements. Altitude and categorical SNAI are intentionally excluded.
RELATIVE_POLARITY = {
    "population_65_plus_pct": "higher_is_worse",
    "income_vulnerability_pct": "higher_is_worse",
    "old_building_stock_pct": "higher_is_worse",
    "heating_degree_days": "higher_is_worse",  # higher potential heating demand
    "dea_access_minutes": "higher_is_worse",
    "upper_secondary_access_minutes": "higher_is_worse",
    "rail_access_minutes": "higher_is_worse",
}


def draw_maps(geojson: dict, destination: Path, *,
              map_mode: str = "absolute", relative_unit: str = "auto") -> list[str]:
    """9 legacy absolute maps by default; optional 7 relative v0.2 previews."""
    if map_mode not in {"absolute", "relative", "both"}:
        raise ValueError("map_mode must be absolute, relative or both")
    names = []
    destination.mkdir(parents=True, exist_ok=True)
    for field, label, unit, filename in MAP_SPECS:
        if map_mode in {"absolute", "both"}:
            render_thematic_map(
                geojson, field, label, unit, destination / filename,
                category_colors=SNAI_COLORS if field == "snai_class" else None,
                footer="EPAH Diagnosis | Sinergie in Canavese | 27 Comuni",
            )
            names.append(f"maps/{filename}")
        if map_mode in {"relative", "both"} and field in RELATIVE_POLARITY:
            relative_filename = "relative_" + filename
            render_thematic_map(
                geojson, field, label + " - differenza dalla media", unit,
                destination / relative_filename, mode="relative",
                direction=RELATIVE_POLARITY[field], relative_unit=relative_unit,
                footer="EPAH Diagnosis | Media non ponderata dei 27 Comuni",
            )
            names.append(f"maps/{relative_filename}")
    return names


def _safe_cell(value):
    """Protect Excel cells from CSV formula injection in untrusted source text."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_workbook(
    path: Path,
    profiles: list[dict],
    comparison: list[dict],
    matrix: list[dict],
    catalog: list[dict],
    registry: dict[str, str],
    readiness: dict,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Overview")
    ws.append(["EPAH Diagnosis v0.1", "Sinergie in Canavese"])
    ws.append(["Scope", "27 municipalities, nine independent core indicators"])
    ws.append(["Observed validated municipality/indicator pairs", 243])
    ws.append(["Model", "Municipal screening, not measured household energy poverty"])
    ws.append(["Composite score", "Not calculated"])
    ws.append(["Release input gate", str(readiness["inputs_ready"])])
    ws.append(["Data reference", "Source year per indicator; see Matrix_Evidence"])
    ws.append(["SNAI", "Categorical A-F, never a numeric ranking"])
    ws = wb.create_sheet("Municipalities")
    header = ["municipality", "istat_code", *NUMERIC_FIELDS, "snai_class"]
    ws.append(header)
    for row in sorted(profiles, key=lambda x: x["municipality"]):
        name = row["municipality"]
        ws.append([name, registry[name], *(float(row[f]) for f in NUMERIC_FIELDS),
                   row["snai_class"]])
    for cells in ws.iter_rows(min_row=2, min_col=3, max_col=10):
        for cell in cells:
            cell.number_format = "0.00"

    ws = wb.create_sheet("Indicators")
    ws.append(["indicator_id", "name", "unit", "source_ids",
               "methodological_maturity", "municipalities", "reference_years"])
    grouped: dict[str, set[str]] = {}
    for row in matrix:
        grouped.setdefault(row["indicator_id"], set()).add(row["year"])
    for item in catalog:
        ws.append([
            item["id"], _safe_cell(item["name"]), item["unit"],
            " | ".join(item.get("source_ids", [])),
            item.get("methodological_maturity", ""),
            27, " | ".join(sorted(grouped.get(item["id"], []))),
        ])
    ws = wb.create_sheet("Comparison")
    columns = list(comparison[0])
    ws.append(columns)
    for item in comparison:
        ws.append([_safe_cell(item.get(key)) for key in columns])

    ws = wb.create_sheet("Matrix_Evidence")
    columns = [
        "municipality", "istat_code", "indicator_id", "status", "value",
        "year", "unit", "source_url", "quality", "limitations",
    ]
    ws.append(columns)
    for row in sorted(matrix, key=lambda r: (r["municipality"], r["indicator_id"])):
        ws.append([
            row["municipality"], registry[row["municipality"]], row["indicator_id"],
            row["status"], _safe_cell(row["value"]), row["year"],
            row["unit"], _safe_cell(row["source_url"]),
            _safe_cell(row["quality"]), _safe_cell(row["limitations"]),
        ])

    for worksheet in wb:
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "C2" if worksheet.title in ("Municipalities", "Matrix_Evidence") else "B2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.row_dimensions[1].height = 30
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="18354D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col in worksheet.columns:
            col_id = get_column_letter(col[0].column)
            maxlen = max((len(str(c.value or "")) for c in col[:60]), default=12)
            worksheet.column_dimensions[col_id].width = min(48, max(14, maxlen + 2))
    wb.save(path)


def _report_footer(canvas, document):
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawString(1.7 * cm, 1.2 * cm,
                      "EPAH Diagnosis v0.1 | FERVORES - Green Community Sinergie in Canavese")
    canvas.drawRightString(19.3 * cm, 1.2 * cm, str(document.page))


def write_report(path: Path, profiles: list[dict], catalog: list[dict], readiness: dict) -> None:
    """Methodological PDF report: descriptive aggregates, no fictitious causation."""
    doc = SimpleDocTemplate(
        str(path), pagesize=(21 * cm, 29.7 * cm),
        rightMargin=1.7 * cm, leftMargin=1.7 * cm,
        topMargin=1.9 * cm, bottomMargin=2 * cm,
        title="EPAH Diagnosis v0.1 - Diagnosi territoriale Canavese",
        author="EPAH Diagnosis - FERVORES",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="EpaTitle", parent=styles["Title"], fontSize=19, leading=24,
        textColor=colors.HexColor("#163650"), spaceAfter=13,
    ))
    styles.add(ParagraphStyle(
        name="EpaBody", parent=styles["BodyText"], fontSize=9.5,
        leading=14.2, spaceAfter=10,
    ))
    styles.add(ParagraphStyle(
        name="EpaSmall", parent=styles["BodyText"], fontSize=8,
        leading=11, spaceAfter=6,
    ))
    story = [
        Paragraph("Diagnosi territoriale della vulnerabilita energetica",
                  styles["EpaTitle"]),
        Paragraph(
            "<b>Green Community Sinergie in Canavese - 27 Comuni</b><br/>"
            "EPAH Diagnosis | Pacchetto v0.1 | Sintesi tecnica per CERS FERVORES",
            styles["EpaBody"],
        ),
        Paragraph(
            "Questo documento e uno <b>screening territoriale esplorativo</b> "
            "basato su indicatori comunali aggregati. Non individua famiglie "
            "in poverta energetica, non misura il consumo energetico individuale "
            "e non definisce un indice composito o una graduatoria complessiva.",
            styles["EpaBody"],
        ),
        Paragraph("Impostazione e risultati", styles["Heading2"]),
        Paragraph(
            f"Sono inclusi <b>{len(profiles)} Comuni e {len(catalog)} dimensioni</b>. "
            "Ogni indicatore mantiene il proprio anno di riferimento. "
            "Le variabili demografiche, reddituali, edilizie, climatiche e "
            "di accessibilita descrivono aspetti differenti della vulnerabilita, "
            "senza implicare causalita individuale o equivalenza fra dimensioni.",
            styles["EpaBody"],
        ),
        Paragraph("Indicatori numerici - variabilita tra Comuni", styles["Heading2"]),
    ]
    table_rows = [["Indicatore", "Min", "Mediana", "Max"]]
    labels = {
        "population_65_plus_pct": "Popolazione 65+ (%)",
        "income_vulnerability_pct": "Reddito basso (%)",
        "old_building_stock_pct": "Edifici ante 1981 (%)",
        "heating_degree_days": "Gradi-giorno (GG)",
        "mean_altitude_m": "Altitudine media (m)",
        "dea_access_minutes": "Accesso DEA (min)",
        "upper_secondary_access_minutes": "Scuola superiore (min)",
        "rail_access_minutes": "Ferrovia (min)",
    }
    for field in NUMERIC_FIELDS:
        values = sorted(float(p[field]) for p in profiles)
        center = statistics_median(values)
        table_rows.append([
            labels[field], f"{values[0]:.1f}", f"{center:.1f}", f"{values[-1]:.1f}"
        ])
    table = Table(table_rows, colWidths=[8.1 * cm, 2.7 * cm, 2.7 * cm, 2.7 * cm],
                  repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#18354D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
            colors.whitesmoke, colors.white,
        ]),
        ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#CBD5E1")),
    ]))
    story.extend([table, Spacer(1, 0.5 * cm)])
    classes = Counter(p["snai_class"] for p in profiles)
    story.append(Paragraph(
        "Classificazione SNAI 2020 (categorie, non punteggi): "
        + ", ".join(f"<b>{c}</b>: {classes.get(c, 0)}" for c in "ABCDEF") + ".",
        styles["EpaBody"],
    ))
    story.append(Paragraph(
        "Le statistiche di sintesi qui riportate confrontano Comuni e non "
        "sono medie ponderate per la popolazione. Interpretare la dispersione "
        "insieme alle mappe tematiche e alle note sui dati originali.",
        styles["EpaSmall"],
    ))
    story.append(PageBreak())
    story.extend([
        Paragraph("Fonti, periodi e limiti metodologici", styles["EpaTitle"]),
        Paragraph(
            "Le fonti originarie, gli anni di riferimento, le specifiche "
            "limitazioni per Comune e le evidenze di qualita sono riprodotte "
            "nel file Excel (fogli Indicators e Matrix_Evidence). "
            "Le metodologie e le approvazioni umane sono tracciate nel "
            "catalogo e nel release gate.",
            styles["EpaBody"],
        ),
    ])
    sources = []
    for item in catalog:
        years = set()
        if item["id"] in NUMERIC_INDICATORS:
            field = NUMERIC_INDICATORS[item["id"]][1]
        else:
            field = CATEGORICAL_INDICATORS[item["id"]][1]
        years.update(str(p[field]) for p in profiles)
        sources.append([
            Paragraph(escape(str(item["name"])), styles["EpaSmall"]),
            Paragraph(escape(" | ".join(sorted(years))), styles["EpaSmall"]),
            Paragraph(escape(", ".join(item.get("source_ids", [])) or "Da documentare"),
                      styles["EpaSmall"]),
        ])
    table = Table(
        [[Paragraph("Dimensione", styles["EpaSmall"]),
          Paragraph("Anno", styles["EpaSmall"]),
          Paragraph("Fonti registrate", styles["EpaSmall"])], *sources],
        colWidths=[6.3 * cm, 2.6 * cm, 7.3 * cm], repeatRows=1,
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([table, Spacer(1, .4 * cm)])
    story.extend([
        Paragraph("Cautela interpretativa", styles["Heading2"]),
        Paragraph(
            "Il patrimonio edilizio deriva dal censimento 2011 e non misura "
            "l'efficienza energetica degli edifici oggi. I tempi di accesso "
            "ISTAT utilizzano un grafo stradale 2020 e una geografia 2021; "
            "servizi sanitari, scolastici e ferroviari possono avere "
            "riferimenti successivi e perimetri territoriali limitati. "
            "Il dato SNAI e una classificazione storica di policy. "
            "La comparazione non implica una misura diretta di poverta energetica.",
            styles["EpaBody"],
        ),
        Paragraph(
            "Prossimi usi: discussione dei risultati con FERVORES e Comuni; "
            "identificazione delle aree che meritano verifiche di dettaglio; "
            "eventuale raccolta futura di informazioni aggregate aggiuntive "
            "secondo criteri di privacy e affidabilita.",
            styles["EpaBody"],
        ),
        Paragraph(
            "La presenza di questo rapporto indica che gli input della "
            "release sono stati ammessi dal gate strutturale; non sostituisce "
            "la responsabilita del revisore scientifico delle singole fonti.",
            styles["EpaSmall"],
        ),
    ])
    doc.build(story, onFirstPage=_report_footer, onLaterPages=_report_footer)


def statistics_median(sorted_values: list[float]) -> float:
    n = len(sorted_values)
    return sorted_values[n // 2] if n % 2 else (sorted_values[n // 2 - 1] + sorted_values[n // 2]) / 2


def validate_bundle_contents(
    output: Path, registry: dict[str, str],
    required_files: list[str],
) -> None:
    """Validate final structures, not merely file existence."""
    if not required_files or len(required_files) != len(set(required_files)):
        raise ValueError("Missing/duplicate required release artifacts")
    for item in required_files:
        path = (output / item).resolve()
        if not path.is_relative_to(output.resolve()) or not path.is_file() or path.stat().st_size == 0:
            raise ValueError(f"Missing/empty/out-of-root release artifact: {item}")
    for filename in ("municipality_profile.csv", "municipality_comparison.csv"):
        with (output / filename).open(encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        if len(rows) != 27 or {r["municipality"] for r in rows} != set(registry):
            raise ValueError(f"Invalid 27-municipality table: {filename}")
        if filename == "municipality_profile.csv":
            for row in rows:
                for field in NUMERIC_FIELDS:
                    if not math.isfinite(float(row[field])):
                        raise ValueError(f"Nonfinite profile field: {field}")
                if row["snai_class"] not in SNAI_COLORS:
                    raise ValueError("Invalid SNAI class in profile")
    geo = json.loads((output / "municipality_profile.geojson").read_text(encoding="utf-8"))
    check_geojson(geo, registry, enriched=True)
    workbook = load_workbook(output / "completeness.xlsx", read_only=True, data_only=True)
    try:
        wanted = {"Overview", "Municipalities", "Indicators", "Comparison", "Matrix_Evidence"}
        if not wanted <= set(workbook.sheetnames):
            raise ValueError("Diagnostic workbook missing expected sheets")
        if (workbook["Municipalities"].max_row != 28
            or workbook["Indicators"].max_row != 10
            or workbook["Matrix_Evidence"].max_row != 244):
            raise ValueError("Diagnostic workbook coverage is not 27x9")
        observed_codes = {
            str(row[1]).zfill(6) for row in
            list(workbook["Municipalities"].values)[1:]
        }
        if observed_codes != set(registry.values()):
            raise ValueError("Excel municipality codes do not match registry")
    finally:
        workbook.close()
    if (output / "diagnostic_report.pdf").read_bytes()[:5] != b"%PDF-":
        raise ValueError("Diagnostic report is not a PDF")
    for _, _, _, filename in MAP_SPECS:
        raw = (output / "maps" / filename).read_bytes()[:24]
        if len(raw) != 24 or not raw.startswith(b"\x89PNG\r\n\x1a\n"):
            raise ValueError(f"Invalid PNG map: {filename}")
        w, h = struct.unpack(">II", raw[16:24])
        if not 200 <= w <= 8000 or not 200 <= h <= 8000:
            raise ValueError(f"Map PNG has invalid dimensions: {filename}")
    manifest = json.loads((output / "release_manifest.json").read_text(encoding="utf-8"))
    hashes = manifest.get("artifacts", {})
    if set(hashes) != set(required_files) - EXCLUDED_HASH_FILES:
        raise ValueError("Manifest must hash all independent artifacts exactly once")
    for name, item in hashes.items():
        path = output / name
        if item.get("sha256") != sha256_file(path) or item.get("bytes") != path.stat().st_size:
            raise ValueError(f"Manifest checksum does not match: {name}")
    state = json.loads((output / "readiness.json").read_text(encoding="utf-8"))
    if not (state.get("phase") == "release" and state.get("release_ready") is True):
        raise ValueError("Release readiness artifact is not an approved final release state")
